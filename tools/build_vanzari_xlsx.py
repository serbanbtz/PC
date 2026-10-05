#!/usr/bin/env python3
"""Build west-auto/West_Auto_Vanzari.xlsx: a sales log that computes the business analysis
and the missing [X] numbers for the Tesla interview guide.

Usage: python3 tools/build_vanzari_xlsx.py [output.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "west-auto/West_Auto_Vanzari.xlsx"
N = 1000                      # data rows available in the sales log
LAST = N + 1
F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
BOLD = Font(name=F, bold=True)
GREEN = Font(name=F, color="008000")
HEAD = Font(name=F, bold=True, color="FFFFFF")
TITLE = Font(name=F, bold=True, size=14)
YELLOW = PatternFill("solid", start_color="FFFF00")
HEAD_FILL = PatternFill("solid", start_color="1F3A5F")
CALC_FILL = PatternFill("solid", start_color="EEF2F7")
THIN = Side(style="thin", color="C9D1DC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR = '#,##0 "€";(#,##0 "€");-'
PCT = '0.0%;(0.0%);-'
INT = '#,##0;(#,##0);-'
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---------------------------------------------------------------- legend
leg = wb.active
leg.title = "Cum completezi"
rows = [
    ("West Auto · registrul de vânzări", TITLE),
    ("", None),
    ("Ce face fișierul", BOLD),
    ("1. Treci fiecare mașină vândută pe un rând în foaia „Vanzari” (coloanele A–N).", None),
    ("2. Foaia „Analiza” calculează singură vânzările pe ani, marja, zilele în stoc, ponderea EV, finanțarea și recomandările.", None),
    ("3. Foaia „Cifre interviu” scoate cifrele care lipsesc din ghidul Tesla ([X]) și îți scrie propozițiile în engleză.", None),
    ("", None),
    ("Culori", BOLD),
    ("Text albastru pe fond galben = completezi tu. Text negru pe fond gri = formulă, nu o modifica.", None),
    ("", None),
    ("Exemplu de rând completat (nu e în foaia de date, ca să nu strice statisticile):", BOLD),
]
for i, (text, font) in enumerate(rows, start=1):
    c = leg.cell(row=i, column=1, value=text)
    c.font = font or BLACK
headers = ["Data achiziției", "Data vânzării", "Marcă", "Model", "An fabricație", "Combustibil",
           "Km", "Cost achiziție (€)", "Costuri import și pregătire (€)", "Preț vânzare (€)",
           "Finanțare", "Consultant", "Sursă client", "Reclamație"]
example = ["12.02.2025", "03.03.2025", "Tesla", "Model 3", 2021, "Electric", 68000, 21500, 1400,
           25900, "Da", "Andrei", "Recomandare", "Nu"]
for j, (h, v) in enumerate(zip(headers, example), start=1):
    leg.cell(row=12, column=j, value=h).font = BOLD
    leg.cell(row=13, column=j, value=v).font = BLUE
notes = [
    "",
    "Reguli",
    "• Datele se scriu ca dată (ex. 03.03.2025), nu ca text.",
    "• Combustibil, Finanțare, Sursă client și Reclamație au listă: alegi din meniul celulei.",
    "• Costuri import și pregătire = transport, RAR, înmatriculare, reparații, curățenie. Tot ce ai plătit până la vânzare.",
    "• Sumele sunt în euro, fără TVA dacă așa ții evidența. Important e să faci la fel pe toate rândurile.",
    "• Nu ai date pentru toți anii? Completează ce ai. Formulele ignoră rândurile goale.",
    "• Foaia are loc pentru 1.000 de mașini.",
    "",
    "Sursa datelor: completate de tine din evidența West Auto. Nicio cifră din fișier nu e inventată.",
]
for k, text in enumerate(notes, start=14):
    leg.cell(row=k, column=1, value=text).font = BOLD if text == "Reguli" else BLACK
leg.column_dimensions["A"].width = 18
for col in range(2, 15):
    leg.column_dimensions[get_column_letter(col)].width = 16

# ---------------------------------------------------------------- data
ws = wb.create_sheet("Vanzari")
calc_headers = ["Cost total (€)", "Marjă (€)", "Marjă %", "Zile în stoc", "An vânzare", "Electrificat (1/0)"]
for j, h in enumerate(headers + calc_headers, start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center"), BOX
ws.row_dimensions[1].height = 34
widths = [13, 13, 11, 14, 10, 12, 10, 13, 15, 13, 10, 13, 13, 11, 13, 12, 10, 10, 10, 11]
for j, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C2"

fmt_in = {1: "dd.mm.yyyy", 2: "dd.mm.yyyy", 5: "0", 7: INT, 8: EUR, 9: EUR, 10: EUR}
for r in range(2, LAST + 1):
    for j in range(1, 15):
        c = ws.cell(row=r, column=j)
        c.font, c.fill = BLUE, YELLOW
        if j in fmt_in:
            c.number_format = fmt_in[j]
    f = {
        15: f'=IF(J{r}="","",H{r}+I{r})',
        16: f'=IF(J{r}="","",J{r}-O{r})',
        17: f'=IF(OR(J{r}="",J{r}=0),"",P{r}/J{r})',
        18: f'=IF(OR(A{r}="",B{r}=""),"",B{r}-A{r})',
        19: f'=IF(B{r}="","",YEAR(B{r}))',
        20: f'=IF(B{r}="","",IF(OR(F{r}="Electric",F{r}="Plug-in",F{r}="Hibrid"),1,0))',
    }
    for j, formula in f.items():
        c = ws.cell(row=r, column=j, value=formula)
        c.font, c.fill = BLACK, CALC_FILL
    ws.cell(row=r, column=15).number_format = EUR
    ws.cell(row=r, column=16).number_format = EUR
    ws.cell(row=r, column=17).number_format = PCT
    ws.cell(row=r, column=18).number_format = INT
    ws.cell(row=r, column=19).number_format = "0"

for col, options in {"F": "Electric,Plug-in,Hibrid,Benzină,Diesel,GPL",
                     "K": "Da,Nu", "M": "Showroom,Online,Recomandare,Client revenit,Altă sursă",
                     "N": "Da,Nu"}.items():
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
    dv.error, dv.errorTitle = "Alege o valoare din listă.", "Valoare necunoscută"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LAST}")
for col in "AB":
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    dv.error, dv.errorTitle = "Scrie o dată, de exemplu 03.03.2025.", "Dată invalidă"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LAST}")
ws["I1"].comment = Comment("Transport, RAR, înmatriculare, reparații, pregătire. Tot ce ai plătit până la vânzare.", "West Auto")
ws["T1"].comment = Comment("1 dacă mașina e Electric, Plug-in sau Hibrid.", "West Auto")


def rng(col):
    return f"Vanzari!${col}$2:${col}${LAST}"


S, T, J, P, Q, R, F_, K, M, N_, C, D, L = (rng(c) for c in "STJPQRFKMNCDL")

# ---------------------------------------------------------------- analysis
an = wb.create_sheet("Analiza")
an["A1"], an["A1"].font = "Analiza vânzărilor West Auto", TITLE
an["A2"] = "Se calculează singură din foaia „Vanzari”. Galben = poți schimba (ani, modele, consultanți)."
an["A2"].font = Font(name=F, italic=True)


def header_row(sheet, row, labels, start=1):
    for j, h in enumerate(labels, start=start):
        c = sheet.cell(row=row, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center"), BOX
    sheet.row_dimensions[row].height = 34


def put(sheet, row, col, value, fmt=None, font=BLACK, fill=None):
    c = sheet.cell(row=row, column=col, value=value)
    c.font, c.border = font, BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


an["A4"], an["A4"].font = "Pe ani", BOLD
header_row(an, 5, ["An", "Mașini vândute", "Din care electrificate", "% electrificate", "Doar 100% electrice",
                   "Venituri (€)", "Marjă totală (€)", "Marjă medie / mașină (€)", "Marjă %",
                   "Zile medii în stoc", "% cu finanțare", "% din recomandări", "Reclamații"])
for i, year in enumerate(range(2019, 2027)):
    r = 6 + i
    put(an, r, 1, year, "0", BLUE, YELLOW)
    y = f"$A{r}"
    put(an, r, 2, f"=COUNTIFS({S},{y})", INT)
    put(an, r, 3, f"=SUMIFS({T},{S},{y})", INT)
    put(an, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", PCT)
    put(an, r, 5, f'=COUNTIFS({S},{y},{F_},"Electric")', INT)
    put(an, r, 6, f"=SUMIFS({J},{S},{y})", EUR)
    put(an, r, 7, f"=SUMIFS({P},{S},{y})", EUR)
    put(an, r, 8, f"=IF(B{r}=0,0,G{r}/B{r})", EUR)
    put(an, r, 9, f"=IF(F{r}=0,0,G{r}/F{r})", PCT)
    put(an, r, 10, f'=IFERROR(AVERAGEIFS({R},{S},{y}),0)', '0;-0;-')
    put(an, r, 11, f'=IF(B{r}=0,0,COUNTIFS({S},{y},{K},"Da")/B{r})', PCT)
    put(an, r, 12, f'=IF(B{r}=0,0,COUNTIFS({S},{y},{M},"Recomandare")/B{r})', PCT)
    put(an, r, 13, f'=COUNTIFS({S},{y},{N_},"Da")', INT)
tot = 14
put(an, tot, 1, "Total", font=BOLD)
for col in (2, 3, 5, 6, 7, 13):
    L_ = get_column_letter(col)
    put(an, tot, col, f"=SUM({L_}6:{L_}13)", EUR if col in (6, 7) else INT, BOLD)
put(an, tot, 4, f"=IF(B{tot}=0,0,C{tot}/B{tot})", PCT, BOLD)
put(an, tot, 8, f"=IF(B{tot}=0,0,G{tot}/B{tot})", EUR, BOLD)
put(an, tot, 9, f"=IF(F{tot}=0,0,G{tot}/F{tot})", PCT, BOLD)
put(an, tot, 10, f"=IFERROR(AVERAGE({R}),0)", '0;-0;-', BOLD)
put(an, tot, 11, f'=IF(B{tot}=0,0,COUNTIFS({K},"Da")/B{tot})', PCT, BOLD)
put(an, tot, 12, f'=IF(B{tot}=0,0,COUNTIFS({M},"Recomandare")/B{tot})', PCT, BOLD)
an.cell(row=15, column=1, value="Totalul include doar anii din tabel. Schimbă anii din coloana A dacă ai vânzări în afara lor.").font = Font(name=F, italic=True, size=9)

an["A17"], an["A17"].font = "Pe tip de combustibil", BOLD
header_row(an, 18, ["Combustibil", "Mașini vândute", "% din total", "Marjă medie / mașină (€)", "Marjă %", "Zile medii în stoc"])
for i, fuel in enumerate(["Electric", "Plug-in", "Hibrid", "Benzină", "Diesel", "GPL"]):
    r = 19 + i
    put(an, r, 1, fuel)
    put(an, r, 2, f'=COUNTIFS({F_},$A{r},{S},"<>")', INT)
    put(an, r, 3, f"=IF($B${tot}=0,0,B{r}/$B${tot})", PCT)
    put(an, r, 4, f"=IFERROR(AVERAGEIFS({P},{F_},$A{r}),0)", EUR)
    put(an, r, 5, f"=IFERROR(SUMIFS({P},{F_},$A{r})/SUMIFS({J},{F_},$A{r}),0)", PCT)
    put(an, r, 6, f"=IFERROR(AVERAGEIFS({R},{F_},$A{r}),0)", '0;-0;-')

an["A26"], an["A26"].font = "Pe model (scrii tu marca și modelul)", BOLD
header_row(an, 27, ["Marcă", "Model", "Mașini vândute", "Marjă medie / mașină (€)", "Zile medii în stoc", "Preț mediu de vânzare (€)"])
models = [("Tesla", "Model 3"), ("Tesla", "Model Y"), ("Tesla", "Model S"), ("Tesla", "Model X")] + [("", "")] * 8
for i, (mk, md) in enumerate(models):
    r = 28 + i
    put(an, r, 1, mk or None, font=BLUE, fill=YELLOW)
    put(an, r, 2, md or None, font=BLUE, fill=YELLOW)
    cond = f"{C},$A{r},{D},$B{r}"
    put(an, r, 3, f'=IF(A{r}="",0,COUNTIFS({cond}))', INT)
    put(an, r, 4, f'=IF(C{r}=0,0,AVERAGEIFS({P},{cond}))', EUR)
    put(an, r, 5, f'=IFERROR(AVERAGEIFS({R},{cond}),0)', '0;-0;-')
    put(an, r, 6, f'=IF(C{r}=0,0,AVERAGEIFS({J},{cond}))', EUR)

an["H26"], an["H26"].font = "Pe consultant (scrii tu numele)", BOLD
header_row(an, 27, ["Consultant", "Mașini vândute", "Marjă medie / mașină (€)", "% cu finanțare", "Reclamații"], start=8)
for i in range(8):
    r = 28 + i
    put(an, r, 8, None, font=BLUE, fill=YELLOW)
    put(an, r, 9, f'=IF(H{r}="",0,COUNTIFS({L},$H{r},{S},"<>"))', INT)
    put(an, r, 10, f'=IF(I{r}=0,0,AVERAGEIFS({P},{L},$H{r}))', EUR)
    put(an, r, 11, f'=IF(I{r}=0,0,COUNTIFS({L},$H{r},{K},"Da")/I{r})', PCT)
    put(an, r, 12, f'=IF(H{r}="",0,COUNTIFS({L},$H{r},{N_},"Da"))', INT)

for j, w in enumerate([13, 12, 13, 13, 12, 14, 14, 16, 13, 13, 13, 13, 12], start=1):
    an.column_dimensions[get_column_letter(j)].width = w
an.freeze_panes = "B6"

# ---------------------------------------------------------------- interview numbers
iv = wb.create_sheet("Cifre interviu")
iv["A1"], iv["A1"].font = "Cifrele pentru ghidul Tesla", TITLE
iv["A2"] = "Coloana B se calculează din vânzări sau o completezi tu (galben). Coloana D îți dă propoziția în engleză, gata de spus."
iv["A2"].font = Font(name=F, italic=True)

iv["A4"], iv["A4"].font = "Setări", BOLD
settings = [
    ("Anul trecut (ultimul an complet)", 2025, "Pentru „Last year we sold [X] cars”."),
    ("Anul în care ai preluat firma", 2021, "Pentru „from [X] cars a year when I took over”. Schimbă-l cu anul real."),
    ("Câți oameni ai în echipă acum", None, "Nu se poate calcula din vânzări. Completează tu."),
    ("Câți oameni ai angajat de-a lungul timpului", None, "Ajută la întrebarea despre construirea echipei."),
]
for i, (label, val, note) in enumerate(settings):
    r = 5 + i
    put(iv, r, 1, label)
    put(iv, r, 2, val, "0", BLUE, YELLOW)
    iv.cell(row=r, column=3, value=note).font = Font(name=F, italic=True, size=9)

LY, FY, TEAM = "$B$5", "$B$6", "$B$7"
YR = "Analiza!$A$6:$A$13"


def by_year(col, year):
    return f"IFERROR(INDEX(Analiza!${col}$6:${col}$13,MATCH({year},{YR},0)),0)"


iv["A10"], iv["A10"].font = "Cifrele calculate", BOLD
header_row(iv, 11, ["Ce", "Cifra", "Unde intră în ghid", "Ce spui (EN)"])
calc = [
    ("Mașini vândute anul trecut", f"={by_year('B', LY)}", INT, "Q1 Tell me about yourself",
     '="Last year we sold "&TEXT(B12,"0")&" cars."'),
    ("% electrice sau hibride anul trecut", f"={by_year('D', LY)}", PCT, "Q1, Q2",
     '="Last year "&TEXT(B13,"0%")&" of the cars we sold were electric or hybrid."'),
    ("Mașini vândute în anul preluării", f"={by_year('B', FY)}", INT, "Q3 follow-up: next level",
     '=IF(B12>=B14,"We grew from "&TEXT(B14,"0")&" cars a year when I took over to "&TEXT(B12,"0")&" last year.","Sales went from "&TEXT(B14,"0")&" to "&TEXT(B12,"0")&" cars a year. Think about how to frame this before Thursday.")'),
    ("% electrice sau hibride în anul preluării", f"={by_year('D', FY)}", PCT, "Q16 Market shift",
     '="EVs and hybrids went from "&TEXT(B15,"0%")&" of our sales in "&TEXT(' + FY + ',"0")&" to "&TEXT(B13,"0%")&" in "&TEXT(' + LY + ',"0")&"."'),
    ("Creștere vânzări, anul preluării → anul trecut", "=IF(B14=0,0,B12/B14-1)", PCT, "Q3 follow-up",
     '=IF(B14=0,"Fill in sales for the takeover year first.",IF(B16>=0,"That is "&TEXT(B16,"0%")&" growth.","Sales are "&TEXT(-B16,"0%")&" lower than in the takeover year. Do not use this line."))'),
    ("Zile medii în stoc (toți anii)", "=Analiza!$J$14", '0', "Doar dacă te întreabă de operațiuni",
     '="On average a car spent "&TEXT(B17,"0")&" days in stock."'),
    ("% clienți cu finanțare (toți anii)", "=Analiza!$K$14", PCT, "Q4 Underperformer (finanțare)",
     '="About "&TEXT(B18,"0%")&" of our customers bought with financing or leasing."'),
    ("% clienți din recomandări (toți anii)", "=Analiza!$L$14", PCT, "Q12, Q13 (recomandări)",
     '="About "&TEXT(B19,"0%")&" of our customers came through referrals."'),
    ("Reclamații la 100 de mașini (toți anii)", "=IF(Analiza!$B$14=0,0,Analiza!$M$14/Analiza!$B$14*100)", '0.0', "Q11 Difficult customer",
     '="We had "&TEXT(B20,"0.0")&" complaints per 100 cars sold."'),
    ("Mărimea echipei", f'=IF({TEAM}="","",{TEAM})', '0', "Q1, Q3",
     '=IF(B21="","Fill in your team size in B7.","I lead a team of "&TEXT(B21,"0")&" people.")'),
]
for i, (label, formula, fmt, where, say) in enumerate(calc):
    r = 12 + i
    put(iv, r, 1, label)
    put(iv, r, 2, formula, fmt, GREEN if "Analiza!" in formula or "INDEX" in formula else BLACK, CALC_FILL)
    put(iv, r, 3, where)
    put(iv, r, 4, say, font=BLACK).alignment = WRAP

iv["A24"], iv["A24"].font = "Cifre pe care le știi doar tu (completează din memorie, aproximativ e în regulă)", BOLD
header_row(iv, 25, ["Ce", "Cifra", "Unde intră în ghid", "Notă"])
manual = [
    ("Test drive-uri pe lună ale consultantului din povestea cu finanțarea", "Q4", "Aproximativ."),
    ("Conversia lui înainte de coaching (%)", "Q4", "Ex. 15%."),
    ("Conversia lui după coaching (%)", "Q4", ""),
    ("Media de conversie a echipei (%)", "Q4", ""),
    ("În câte săptămâni s-a văzut rezultatul", "Q4", ""),
    ("Livrări în săptămâna cu conflictul", "Q6", ""),
    ("Câte mașini ai avut în luna cu întârzierea la import (vs. plan)", "Q9", "Ex. 6 în loc de 10."),
    ("Câte zile a durat reparația clientului ANPC", "Q11", "Ghidul spune 5. Verifică."),
    ("Câți clienți ți-a recomandat clientul din povestea cu weekendul", "Q12", ""),
    ("Câți din 10 clienți întreabă întâi de electrice", "Q2", "Spune cifra reală, nu 8."),
]
for i, (label, where, note) in enumerate(manual):
    r = 26 + i
    put(iv, r, 1, label).alignment = WRAP
    put(iv, r, 2, None, None, BLUE, YELLOW)
    put(iv, r, 3, where)
    put(iv, r, 4, note).alignment = WRAP
for col, w in zip("ABCD", [46, 12, 26, 62]):
    iv.column_dimensions[col].width = w

wb.calculation.fullCalcOnLoad = True
wb.move_sheet("Cifre interviu", offset=-2)
wb.active = 0
wb.save(OUT)
print("saved", OUT)
